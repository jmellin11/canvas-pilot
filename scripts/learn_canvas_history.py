# SPDX-License-Identifier: AGPL-3.0-or-later
"""Pull the student's Canvas history into context/archive/ (read-only GETs).

For every course the student can still open: syllabus, assignment groups,
assignments + rubrics, their own submissions (grade, late/missing, instructor
comments, rubric scores), and the files they turned in, with extracted text.
Many schools lock concluded courses (HTTP 403), so the past is recovered from
what stays readable: final grades for every enrollment, the student's own
Canvas files (assignment uploads live under `my files/Submissions/<course>/`),
and inbox threads. The `pilot-learn` skill reads this archive to write the
context notes.

    .venv/bin/python scripts/learn_canvas_history.py              # every course
    .venv/bin/python scripts/learn_canvas_history.py --course 69222
    .venv/bin/python scripts/learn_canvas_history.py --no-files
    .venv/bin/python scripts/learn_canvas_history.py --stats-only

Courses already pulled are skipped unless --refresh or --course names them.
Courses whose term has not ended are always re-pulled.
"""
from __future__ import annotations

import argparse
import datetime as dt
import html
import json
import re
import sys
import zipfile
from pathlib import Path
from urllib.parse import urlencode

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src import canvas_client as cc  # noqa: E402

ARCHIVE = ROOT / "context" / "archive"
TEXT_EXTS = {".txt", ".md", ".html", ".htm", ".csv"}


class SignedOut(RuntimeError):
    pass


def _raw_get(url: str) -> tuple[int, dict, bytes]:
    """GET without canvas_client's 401 → interactive re-login. Canvas answers
    401 for access-restricted (concluded) courses too, and those must be
    skipped, not treated as an expired session."""
    be = cc._backend
    if hasattr(be, "_session"):
        r = be._session.get(url, timeout=60)
        return r.status_code, dict(r.headers), r.content
    be._ensure_session()
    r = be._ctx.request.get(url, timeout=60_000)
    return r.status, dict(r.headers), r.body()


def _url(path: str, params: dict | None = None) -> str:
    url = path if path.startswith("http") else f"{cc.BASE}{path}"
    if params:
        url += ("&" if "?" in url else "?") + urlencode(params, doseq=True)
    return url


def fetch_all(path: str, params: dict | None = None) -> list | dict | None:
    """Paginated GET. Returns None when Canvas refuses access."""
    url, out = _url(path, {"per_page": 100, **(params or {})}), []
    while url:
        status, headers, body = _raw_get(url)
        if status in (401, 403, 404):
            if not cc._backend.is_session_alive(timeout_s=10):
                raise SignedOut("Canvas session expired; sign in again and re-run.")
            return None
        if status >= 400:
            raise RuntimeError(f"GET {url} -> HTTP {status}: {body[:300]!r}")
        data = json.loads(body or b"null")
        if not isinstance(data, list):
            return data
        out.extend(data)
        link = headers.get("link") or headers.get("Link") or ""
        url = cc._parse_link_header(link).get("next")
    return out


def html_to_text(raw: str | None) -> str:
    if not raw:
        return ""
    t = re.sub(r"(?is)<(script|style)\b.*?</\1>", "", raw)
    t = re.sub(r"(?i)<br\s*/?>|</(p|div|li|h\d|tr|blockquote)>", "\n", t)
    t = re.sub(r"(?i)<li\b[^>]*>", "- ", t)
    t = html.unescape(re.sub(r"<[^>]+>", "", t)).replace("\xa0", " ")
    t = re.sub(r"[ \t]+", " ", t)
    return re.sub(r"\n\s*\n+", "\n\n", t).strip()


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    try:
        if ext == ".docx":
            import docx

            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            return "\n".join(parts).strip()
        if ext == ".pdf":
            import pymupdf

            with pymupdf.open(str(path)) as pdf:
                return "\n".join(page.get_text() for page in pdf).strip()
        if ext == ".pptx":
            with zipfile.ZipFile(path) as z:
                slides = sorted(
                    (n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                    key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[1]).group(1)),
                )
                return "\n\n".join(
                    " ".join(html.unescape(t) for t in re.findall(r"<a:t>(.*?)</a:t>", z.read(n).decode("utf-8", "ignore")))
                    for n in slides
                ).strip()
        if ext in TEXT_EXTS:
            raw = path.read_text(encoding="utf-8", errors="ignore")
            return html_to_text(raw) if ext in (".html", ".htm") else raw.strip()
    except Exception as e:  # corrupt or password-protected file
        return f"[text extraction failed: {type(e).__name__}]"
    return ""


def safe_name(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("_")[:80] or "file"


def list_courses() -> list[dict]:
    return fetch_all("/courses", {
        "include[]": ["term", "total_scores", "teachers"],
        "state[]": ["available", "completed"],
    }) or []


def course_meta(course: dict) -> dict:
    term = course.get("term") or {}
    enrollment = next((e for e in course.get("enrollments") or [] if e.get("type") == "student"), {})
    return {
        "id": course["id"],
        "name": course.get("name"),
        "code": course.get("course_code"),
        "term": term.get("name"),
        "term_end": term.get("end_at") or course.get("end_at"),
        "start_at": course.get("start_at"),
        "teachers": [t.get("display_name") for t in course.get("teachers") or []],
        "final_score": enrollment.get("computed_final_score"),
        "final_grade": enrollment.get("computed_final_grade"),
        "restricted": bool(course.get("access_restricted_by_date")),
    }


def is_current(meta: dict) -> bool:
    end = meta.get("term_end")
    if not end:
        return False
    return dt.datetime.fromisoformat(end.replace("Z", "+00:00")) > dt.datetime.now(dt.timezone.utc)


def pull_course(meta: dict, me: int, with_files: bool) -> dict:
    cid = meta["id"]
    out_dir = ARCHIVE / str(cid)
    out_dir.mkdir(parents=True, exist_ok=True)

    detail = fetch_all(f"/courses/{cid}", {"include[]": ["syllabus_body"]}) or {}
    groups = fetch_all(f"/courses/{cid}/assignment_groups") or []
    assignments = fetch_all(f"/courses/{cid}/assignments")
    if assignments is None:
        meta = {**meta, "status": "no_access"}
        (out_dir / "course.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")
        return meta
    subs = fetch_all(f"/courses/{cid}/students/submissions", {
        "student_ids[]": [me],
        "include[]": ["submission_comments", "rubric_assessment"],
    }) or []
    sub_by_aid = {s["assignment_id"]: s for s in subs}
    group_names = {g["id"]: g.get("name") for g in groups}

    records, n_files = [], 0
    for a in assignments:
        rubric = a.get("rubric") or []
        crit_name = {c["id"]: c.get("description") for c in rubric}
        s = sub_by_aid.get(a["id"]) or {}
        texts, files = [], []
        if s.get("body"):
            texts.append(html_to_text(s["body"]))
        for entry in s.get("discussion_entries") or []:
            if entry.get("user_id") == me and entry.get("message"):
                texts.append(html_to_text(entry["message"]))
        for att in s.get("attachments") or []:
            fname = f"{a['id']}__{safe_name(att.get('display_name') or att.get('filename') or 'file')}"
            dest = out_dir / "submitted" / fname
            if with_files and att.get("url") and not dest.exists():
                status, _, body = _raw_get(att["url"])
                if status == 200:
                    dest.parent.mkdir(parents=True, exist_ok=True)
                    dest.write_bytes(body)
            text = extract_text(dest) if dest.exists() else ""
            if text:
                texts.append(f"[file: {att.get('display_name')}]\n{text}")
            files.append({"name": att.get("display_name"), "path": str(dest.relative_to(ROOT)) if dest.exists() else None,
                          "size": att.get("size")})
            n_files += dest.exists()
        text_path = None
        if texts:
            text_path = out_dir / "text" / f"{a['id']}.txt"
            text_path.parent.mkdir(parents=True, exist_ok=True)
            text_path.write_text("\n\n".join(texts), encoding="utf-8")
        records.append({
            "id": a["id"],
            "name": a.get("name"),
            "group": group_names.get(a.get("assignment_group_id")),
            "types": a.get("submission_types"),
            "points": a.get("points_possible"),
            "due_at": a.get("due_at"),
            "url": a.get("html_url"),
            "description": html_to_text(a.get("description")),
            "rubric": [{"criterion": c.get("description"), "details": c.get("long_description"),
                        "points": c.get("points")} for c in rubric],
            "submission": {
                "state": s.get("workflow_state"),
                "submitted_at": s.get("submitted_at"),
                "late": s.get("late"),
                "missing": s.get("missing"),
                "excused": s.get("excused"),
                "score": s.get("score"),
                "grade": s.get("grade"),
                "attempt": s.get("attempt"),
                "comments": [{"author": c.get("author_name"), "is_me": c.get("author_id") == me,
                              "text": c.get("comment"), "at": c.get("created_at")}
                             for c in s.get("submission_comments") or []],
                "rubric_scores": [{"criterion": crit_name.get(k, k), "points": v.get("points"),
                                   "comment": v.get("comments")}
                                  for k, v in (s.get("rubric_assessment") or {}).items()],
                "files": files,
                "text_path": str(text_path.relative_to(ROOT)) if text_path else None,
            },
        })

    meta = {**meta, "status": "ok", "syllabus": html_to_text(detail.get("syllabus_body")),
            "groups": [{"name": g.get("name"), "weight": g.get("group_weight")} for g in groups],
            "pulled_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")}
    (out_dir / "course.json").write_text(json.dumps(meta, indent=2, ensure_ascii=False), encoding="utf-8")
    (out_dir / "assignments.json").write_text(json.dumps(records, indent=2, ensure_ascii=False), encoding="utf-8")
    return {**meta, "assignments": len(records), "files": n_files}


def pull_grades(course_names: dict[int, str]) -> list[dict]:
    """Final grades for every enrollment. Concluded courses stay listed here
    even when the course itself is access-restricted."""
    enrollments = fetch_all("/users/self/enrollments", {
        "state[]": ["active", "completed", "inactive"], "type[]": ["StudentEnrollment"],
    }) or []
    grades = [{
        "course_id": e["course_id"],
        "name": course_names.get(e["course_id"]),
        "state": e.get("enrollment_state"),
        # current_score ignores ungraded work; final_score counts it as zero.
        "current_score": (e.get("grades") or {}).get("current_score"),
        "final_score": (e.get("grades") or {}).get("final_score"),
        "final_grade": (e.get("grades") or {}).get("final_grade"),
    } for e in enrollments]
    (ARCHIVE / "grades.json").write_text(json.dumps(grades, indent=2, ensure_ascii=False), encoding="utf-8")
    return grades


def pull_personal_files(with_files: bool) -> list[dict]:
    """The student's own Canvas files. Uploads to assignments land in
    `my files/Submissions/<course>/`, which stays readable after a course
    concludes, so this is the main record of past work."""
    folders = {f["id"]: f.get("full_name") or "" for f in fetch_all("/users/self/folders") or []}
    out = []
    for f in fetch_all("/users/self/files") or []:
        folder = folders.get(f.get("folder_id"), "").removeprefix("my files").strip("/")
        name = f.get("display_name") or f.get("filename") or "file"
        record = {"id": f["id"], "name": name, "folder": folder, "course": folder.removeprefix("Submissions/") or None,
                  "created_at": f.get("created_at"), "size": f.get("size"), "kind": f.get("mime_class"),
                  "path": None, "text_path": None, "words": 0}
        if f.get("mime_class") in ("image", "video", "audio"):
            out.append(record)
            continue
        dest = ARCHIVE / "files" / safe_name(folder or "root") / f"{f['id']}__{safe_name(name)}"
        if with_files and f.get("url") and not dest.exists():
            status, _, body = _raw_get(f["url"])
            if status == 200:
                dest.parent.mkdir(parents=True, exist_ok=True)
                dest.write_bytes(body)
        if dest.exists():
            record["path"] = str(dest.relative_to(ROOT))
            text = extract_text(dest)
            if text:
                text_path = dest.with_name(dest.name + ".txt")
                text_path.write_text(text, encoding="utf-8")
                record["text_path"] = str(text_path.relative_to(ROOT))
                record["words"] = len(text.split())
        out.append(record)
    (ARCHIVE / "files.json").write_text(json.dumps(out, indent=2, ensure_ascii=False), encoding="utf-8")
    return out


def pull_inbox(me: int) -> list[dict]:
    """Canvas inbox threads (instructor messages, group work), all scopes."""
    seen, out = set(), []
    for scope in ("inbox", "sent", "archived"):
        for c in fetch_all("/conversations", {"scope": scope}) or []:
            if c["id"] in seen:
                continue
            seen.add(c["id"])
            full = fetch_all(f"/conversations/{c['id']}") or {}
            people = {p["id"]: p.get("name") for p in full.get("participants") or []}
            out.append({
                "id": c["id"], "subject": c.get("subject"), "course": c.get("context_name"),
                "course_id": int(c["context_code"].split("_", 1)[1])
                if (c.get("context_code") or "").startswith("course_") else None,
                "messages": [{"author": people.get(m.get("author_id")), "is_me": m.get("author_id") == me,
                              "at": m.get("created_at"), "body": (m.get("body") or "").strip()}
                             for m in reversed(full.get("messages") or [])],
            })
    (ARCHIVE / "inbox.json").write_text(json.dumps(out, indent=2, ensure_ascii=False), encoding="utf-8")
    return out


def pct(a: dict) -> float | None:
    s, p = a["submission"].get("score"), a.get("points")
    return round(100 * s / p, 1) if s is not None and p else None


def write_stats() -> dict:
    courses, rows = [], []
    for course_json in sorted(ARCHIVE.glob("*/course.json")):
        meta = json.loads(course_json.read_text(encoding="utf-8"))
        a_path = course_json.parent / "assignments.json"
        items = json.loads(a_path.read_text(encoding="utf-8")) if a_path.exists() else []
        for a in items:
            rows.append((meta, a))
        graded = [p for p in (pct(a) for a in items) if p is not None]
        courses.append({
            "id": meta["id"], "name": meta.get("name"), "term": meta.get("term"), "status": meta.get("status"),
            "teachers": meta.get("teachers"), "final_score": meta.get("final_score"),
            "final_grade": meta.get("final_grade"), "assignments": len(items),
            "graded": len(graded), "avg_pct": round(sum(graded) / len(graded), 1) if graded else None,
            "late": sum(bool(a["submission"].get("late")) for a in items),
            "missing": sum(bool(a["submission"].get("missing")) for a in items),
            "with_instructor_comments": sum(any(not c["is_me"] for c in a["submission"]["comments"]) for a in items),
            "with_text": sum(bool(a["submission"].get("text_path")) for a in items),
        })

    by_type: dict[str, list[float]] = {}
    for _, a in rows:
        p = pct(a)
        if p is not None:
            by_type.setdefault(",".join(a.get("types") or ["none"]), []).append(p)
    lowest = sorted(((pct(a), m, a) for m, a in rows if pct(a) is not None and (a.get("points") or 0) >= 5),
                    key=lambda t: t[0])[:30]
    def load(name: str) -> list:
        p = ARCHIVE / name
        return json.loads(p.read_text(encoding="utf-8")) if p.exists() else []

    files, grades, inbox = load("files.json"), load("grades.json"), load("inbox.json")
    by_course: dict[str, dict] = {}
    for f in files:
        if f.get("text_path"):
            row = by_course.setdefault(f.get("course") or "(not in a course folder)", {"files": 0, "words": 0})
            row["files"] += 1
            row["words"] += f.get("words") or 0
    scores = [g["current_score"] for g in grades if g.get("current_score") is not None]
    stats = {
        "generated_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "enrollment_grades": {"n": len(grades), "with_score": len(scores),
                              "scores": sorted(scores, reverse=True)},
        "own_files_with_text": by_course,
        "inbox": {"threads": len(inbox), "by_course": dict(sorted(
            ((k, sum(1 for c in inbox if c.get("course") == k)) for k in {c.get("course") for c in inbox}),
            key=lambda kv: -kv[1]))},
        "courses": courses,
        "by_submission_type": {k: {"n": len(v), "avg_pct": round(sum(v) / len(v), 1)} for k, v in by_type.items()},
        "lowest_scored": [{"pct": p, "course": m.get("name"), "assignment": a.get("name"), "points": a.get("points"),
                           "comments": [c["text"][:300] for c in a["submission"]["comments"] if not c["is_me"]][:2]}
                          for p, m, a in lowest],
    }
    (ARCHIVE / "stats.json").write_text(json.dumps(stats, indent=2, ensure_ascii=False), encoding="utf-8")
    return stats


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--course", type=int, action="append", help="only this course id (repeatable)")
    ap.add_argument("--refresh", action="store_true", help="re-pull courses already in the archive")
    ap.add_argument("--no-files", action="store_true", help="skip downloading submitted files")
    ap.add_argument("--stats-only", action="store_true", help="just rebuild stats.json from the archive")
    ap.add_argument("--only", choices=["courses", "grades", "files", "inbox"], action="append",
                    help="pull just these parts (repeatable); default is all")
    args = ap.parse_args()
    parts = set(args.only or ["courses", "grades", "files", "inbox"])
    if args.course:
        parts = {"courses"}

    ARCHIVE.mkdir(parents=True, exist_ok=True)
    if not args.stats_only:
        me = fetch_all("/users/self")
        if not me:
            raise SignedOut("Canvas session expired; sign in again and re-run.")
        metas = [course_meta(c) for c in list_courses()]
        if "inbox" in parts:
            i = pull_inbox(me["id"])
            print(f"inbox: {len(i)} threads", flush=True)
        if "grades" in parts:
            # Locked courses hide their names; inbox threads still carry them.
            inbox_path = ARCHIVE / "inbox.json"
            names = {t["course_id"]: t["course"] for t in json.loads(inbox_path.read_text(encoding="utf-8"))
                     if t.get("course_id")} if inbox_path.exists() else {}
            names.update({m["id"]: m["name"] for m in metas if m.get("name")})
            g = pull_grades(names)
            print(f"grades: {len(g)} enrollments, {sum(bool(x['name']) for x in g)} named", flush=True)
        if "files" in parts:
            f = pull_personal_files(with_files=not args.no_files)
            print(f"files: {len(f)} personal files, {sum(bool(x['text_path']) for x in f)} with text", flush=True)
        if "courses" not in parts:
            metas = []
        if args.course:
            metas = [m for m in metas if m["id"] in set(args.course)]
        index = []
        for m in metas:
            done = (ARCHIVE / str(m["id"]) / "course.json").exists()
            if m["restricted"]:
                index.append({**m, "status": "restricted"})
                print(f"skip  {m['id']}: access restricted (concluded course)", flush=True)
                continue
            if done and not (args.refresh or args.course or is_current(m)):
                index.append({**json.loads((ARCHIVE / str(m["id"]) / "course.json").read_text(encoding="utf-8")),
                              "status": "cached"})
                print(f"cache {m['id']}: {m['name']}", flush=True)
                continue
            r = pull_course(m, me["id"], with_files=not args.no_files)
            index.append({k: v for k, v in r.items() if k != "syllabus"})
            print(f"pull  {m['id']}: {m['name']} [{r['status']}] "
                  f"{r.get('assignments', 0)} assignments, {r.get('files', 0)} files", flush=True)
        if "courses" in parts and not args.course:
            (ARCHIVE / "index.json").write_text(json.dumps(index, indent=2, ensure_ascii=False), encoding="utf-8")

    stats = write_stats()
    print(f"stats: {len(stats['courses'])} courses in archive -> {(ARCHIVE / 'stats.json').relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
